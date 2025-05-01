from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

# --- 1) 워크북/워크시트 생성 ---
wb = Workbook()
ws = wb.active
ws.title = "견적서"

# --- 2) 기본 스타일 세팅 ---
thin_border = Border(
    left=Side(style="thin"), 
    right=Side(style="thin"), 
    top=Side(style="thin"), 
    bottom=Side(style="thin")
)
medium_border = Border(
    left=Side(style="medium"), 
    right=Side(style="medium"), 
    top=Side(style="medium"), 
    bottom=Side(style="medium")
)
center_align = Alignment(horizontal="center", vertical="center")
left_align = Alignment(horizontal="left", vertical="center")
right_align = Alignment(horizontal="right", vertical="center")

title_font = Font(name="맑은 고딕", size=16, bold=True)
header_font = Font(name="맑은 고딕", size=12, bold=True)
bold_font = Font(name="맑은 고딕", size=11, bold=True)
normal_font = Font(name="맑은 고딕", size=11)
small_font = Font(name="맑은 고딕", size=10)

# --- 3) 시트 전체 기본 설정 ---
# 열 너비 설정 (PDF 비율에 맞춰 대략 조정)
col_widths = {
    1: 20,  # A
    2: 14,  # B
    3: 10,  # C
    4: 12,  # D
    5: 18   # E
}
for col_idx, width in col_widths.items():
    ws.column_dimensions[get_column_letter(col_idx)].width = width

# 행 높이도 일부 조정
row_heights = {
    1: 30,  # 견적서 제목
    2: 20,
    3: 20,
    4: 20,
    5: 20,
    6: 30,  # 견적금액 표기
    8: 25,  # 테이블 헤더
}
for row_idx, height in row_heights.items():
    ws.row_dimensions[row_idx].height = height

# --- 4) 상단 '견 적 서' 제목 ---
ws.merge_cells("A1:E1")
ws["A1"] = "견 적 서"
ws["A1"].font = title_font
ws["A1"].alignment = center_align

# --- 5) 좌측 고객 정보 ---
ws["A2"] = "배세원 귀하"
ws["A2"].font = bold_font

ws["A3"] = "견적일: 2024년 4월 29일"
ws["A4"] = "유효기간: 견적일로부터 1개월"
ws["A5"] = "전화: 070-7099-1011~5 / 팩스: 070-7159-1514"

# --- 6) 우측 공급자 정보(대표자, 등록번호, 주소 등) ---
ws.merge_cells("C2:E2")
ws["C2"] = "공 급 자: 티파니닷컴 신일식"
ws["C2"].font = bold_font
ws["C2"].alignment = right_align

ws.merge_cells("C3:E3")
ws["C3"] = "등록번호: 203-43-69231"
ws["C3"].alignment = right_align

ws.merge_cells("C4:E4")
ws["C4"] = "주소: 서울시 광진구 중곡동 긴고랑로77 지하1층"
ws["C4"].alignment = right_align

# --- 7) 견적금액(공급가액 + 세액) 표시 ---
# PDF 상단에 굵게 표시된 합계 금액처럼 보이도록 구성
ws.merge_cells("A6:D6")
ws["A6"] = "견적금액 (공급가액 + 세액)"
ws["A6"].font = bold_font
ws["A6"].alignment = right_align

ws["E6"] = "₩1,848,000"
ws["E6"].font = Font(name="맑은 고딕", size=12, bold=True, color="000000")
ws["E6"].alignment = right_align

# --- 8) 테이블 헤더 (row=8) ---
table_headers = ["품명", "규격", "수량", "단가", "공급가액"]
for col_idx, header in enumerate(table_headers, start=1):
    cell = ws.cell(row=8, column=col_idx, value=header)
    cell.font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
    cell.alignment = center_align
    cell.fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    cell.border = thin_border

# --- 9) 테이블 본문 데이터 ---
table_data = [
    ["야구잠바 기본", "PCS", 30, 51000, 1530000],
    ["등판 대형 자수", "PCS", 30, 3000, 90000],
    ["프리미엄 습식레자", "PCS", 30, 2000, 60000],
    ["이니셜 서비스", "", "", "", 0],
]
start_row = 9
for i, row_data in enumerate(table_data):
    for j, value in enumerate(row_data):
        cell = ws.cell(row=start_row + i, column=j + 1, value=value)
        cell.font = normal_font
        cell.alignment = center_align if j >= 1 else left_align
        cell.border = thin_border

# --- 10) 합계/부가세/총합 등 추가 ---
# 합계
sum_row = start_row + len(table_data)
ws.cell(row=sum_row, column=1, value="합계").font = bold_font
ws.cell(row=sum_row, column=5, value=1680000).font = bold_font
for col in range(1, 6):
    cell = ws.cell(row=sum_row, column=col)
    cell.border = thin_border
    if col != 1 and col != 5:
        cell.value = ""

# VAT
vat_row = sum_row + 1
ws.cell(row=vat_row, column=1, value="VAT").font = bold_font
ws.cell(row=vat_row, column=5, value=168000).font = bold_font
for col in range(1, 6):
    cell = ws.cell(row=vat_row, column=col)
    cell.border = thin_border
    if col != 1 and col != 5:
        cell.value = ""

# --- 11) 테이블 전체에 굵은 외곽선(중간 border는 얇게) ---
# 테이블 범위: A8:E(마지막 행)
last_row = vat_row
for row_idx in range(8, last_row + 1):
    for col_idx in range(1, 6):
        cell = ws.cell(row=row_idx, column=col_idx)
        b = cell.border
        # 상단 테두리
        if row_idx == 8:
            top_side = Side(style="medium")
        else:
            top_side = Side(style=b.top.style if b.top else "thin")
        # 하단 테두리
        if row_idx == last_row:
            bottom_side = Side(style="medium")
        else:
            bottom_side = Side(style=b.bottom.style if b.bottom else "thin")
        # 좌측 테두리
        if col_idx == 1:
            left_side = Side(style="medium")
        else:
            left_side = Side(style=b.left.style if b.left else "thin")
        # 우측 테두리
        if col_idx == 5:
            right_side = Side(style="medium")
        else:
            right_side = Side(style=b.right.style if b.right else "thin")

        cell.border = Border(
            top=top_side, bottom=bottom_side,
            left=left_side, right=right_side
        )

# --- 12) 파일 저장 ---
output_path = "티파니닷컴_견적_레이아웃.xlsx"  # 현재 디렉토리에 저장
wb.save(output_path)

print(f"생성된 엑셀 파일: {output_path}")
